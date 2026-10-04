"""Точка входа: сбор Apify Store -> MySQL -> ниши, услуги, топ-3 в каждой услуге."""
from __future__ import annotations

import argparse
import json
import datetime as dt
import logging
import sys
import time
from pathlib import Path

import pandas as pd

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    __package__ = "tracker"

from . import catalog, content, db, ids, monitor
from .apify_api import API_URL, DEFAULT_MAX_ACTORS, Collector, normalize
from .config import ScoreConfig, parse_weights
from .scoring import build_rankings, formula_text

log = logging.getLogger("apify-tracker")


def make_cfg(args) -> ScoreConfig:
    w = parse_weights(args.weights)
    return ScoreConfig(per_service=args.per_service, min_niche_actors=args.min_niche_actors,
                       min_service_actors=args.min_service_actors, min_service_pool=args.min_service_pool,
                       min_discovered_service=args.min_discovered_service, auto_services=not args.no_auto_services,
                       min_users=args.min_users, compute_per_1000=args.compute_cost,
                       monthly_results=args.monthly_results, bayes_prior=args.bayes_prior,
                       w_price=w[0], w_users=w[1], w_rating=w[2], w_reviews=w[3], niches_file=args.niches_file)


def make_rewrite_cfg(args) -> content.RewriteConfig:
    return content.RewriteConfig(
        provider=args.rewrite_provider, model=args.rewrite_model or content.RewriteConfig().model, lang=args.rewrite_lang,
        sources_per_service=args.rewrite_sources, details_ttl_days=args.details_ttl_days,
        fetch_details=not args.no_fetch_details, workers=args.rewrite_workers,
        window_s=args.rewrite_window_hours * 3600, flush_s=args.flush_minutes * 60,
        min_age_days=args.rewrite_min_age_days, max_similarity=args.rewrite_max_sim,
        niche=args.rewrite_niche, limit=args.rewrite_limit, scope=args.catalog_scope, min_users=args.min_users, force=args.rewrite_force, delay=args.delay)


def score_and_publish(conn, run_id: int, snapshot: pd.DataFrame, cfg: ScoreConfig) -> None:
    rk = build_rankings(snapshot, cfg)
    db.persist_rankings(conn, run_id, rk, cfg)


def collect_rows(args) -> tuple[list[dict], dict, int, bool]:
    """Обход Store (или демо) -> (строки Actor'ов, категории, всего в Store, покрытие достаточное)."""
    started = time.time()
    coverage_ok = True
    if args.demo:
        from .demo_data import make_items
        items, total = make_items(args.demo), args.demo
        log.info("Демо-режим: сгенерировано %s Actor'ов (в Apify запросы не идут)", len(items))
    else:
        col = Collector(args.api_url, args.delay, workers=args.workers, max_actors=args.max_actors,
                        include_unrunnable=not args.exclude_unrunnable, max_terms=args.max_terms)
        items, total = col.collect(args.mode)
        col.enrich_pricing(items, args.pricing_details)
        goal = min(total, args.max_actors) if (total and args.max_actors) else (total or args.max_actors)
        coverage = len(items) / goal * 100 if goal else 0
        log.info("Собрано %s (цель/всего в Store: %s, %.1f%%), запросов: %s, неудачных: %s, %.0f сек",
                 len(items), goal, coverage, col.requests_made, col.failed, time.time() - started)
        if goal and coverage < 90:
            coverage_ok = False
            log.warning("Покрытие ниже 90%%: используйте --mode full, увеличьте --max-terms или проверьте 429.")
    if not items:
        raise RuntimeError("API вернул 0 Actor'ов, снимок не сохранён")

    rows, cats = {}, {}
    for it in items:
        row, cs = normalize(it)
        rows[row["actor_key"]] = row          # дубли по username/name схлопываются
        cats[row["actor_key"]] = cs
    return list(rows.values()), cats, total, coverage_ok


def run_once(args) -> None:
    started = time.time()
    cfg = make_cfg(args)
    rows, cats, total, coverage_ok = collect_rows(args)

    conn = db.connect(db.dsn_from_env(args.dsn))
    try:
        db.init_schema(conn)
        ids.resolve_ids(conn, rows, workers=args.workers, delay=args.delay, use_site=not args.no_site_ids)
        today = dt.date.today().isoformat()
        run_id = db.persist_snapshot(conn, today, rows, cats, total)
        score_and_publish(conn, run_id, pd.DataFrame(rows), cfg)
        db.prune(conn, args.keep_days)
        try:   # таблица 1: цены всех акторов (ошибка не откатывает уже опубликованные рейтинги)
            monitor.sync(conn, rows, cfg, mark_inactive=coverage_ok)
        except Exception:
            log.exception("Мониторинг цен завершился ошибкой (снимок и рейтинги уже опубликованы)")
        if args.xlsx:
            export_xlsx(conn, args.xlsx, cfg)
        if args.with_catalog:   # таблица 2: рерайт только новых ID
            try:
                catalog.run_catalog(conn, make_rewrite_cfg(args))
            except Exception:
                log.exception("Каталог завершился ошибкой (сбор и рейтинги уже опубликованы)")
        if args.with_rewrite:   # рерайт услуг (ниша x услуга) идёт уже по опубликованному снимку
            try:
                content.run_rewrite(conn, make_rewrite_cfg(args))
            except Exception:
                log.exception("Рерайт завершился ошибкой (сбор и рейтинги уже опубликованы)")
    finally:
        conn.close()
    log.info("Готово за %.0f сек", time.time() - started)


def monitor_once(args) -> None:
    """Один проход мониторинга цен: обход Store -> ID -> actor_price_monitor. Без снимка, рейтингов и рерайта."""
    started = time.time()
    cfg = make_cfg(args)
    rows, _, _, coverage_ok = collect_rows(args)
    conn = db.connect(db.dsn_from_env(args.dsn))
    try:
        db.init_schema(conn)
        ids.resolve_ids(conn, rows, workers=args.workers, delay=args.delay, use_site=not args.no_site_ids)
        monitor.sync(conn, rows, cfg, mark_inactive=coverage_ok)
    finally:
        conn.close()
    log.info("Мониторинг цен: проход за %.0f мин", (time.time() - started) / 60)


def monitor_loop(args) -> None:
    """Проход мониторинга раз в --monitor-loop-hours часов (отсчёт от начала прохода)."""
    period = args.monitor_loop_hours * 3600
    log.info("Мониторинг цен: цикл раз в %.1f ч", args.monitor_loop_hours)
    while True:
        t = time.time()
        try:
            monitor_once(args)
        except Exception:
            log.exception("Проход мониторинга завершился ошибкой, повторим по расписанию")
        time.sleep(max(60.0, period - (time.time() - t)))


def catalog_cmd(args) -> None:
    """Одна итерация каталога: рерайт новых ID (окно --rewrite-window-hours)."""
    conn = db.connect(db.dsn_from_env(args.dsn))
    try:
        db.init_schema(conn)
        catalog.run_catalog(conn, make_rewrite_cfg(args))
    finally:
        conn.close()


def probe_api_cmd(args) -> None:
    """Реальные лимиты Store API на вашей сети: сколько записей отдаётся за запрос и работает ли offset."""
    import requests
    from .apify_api import get_json
    s = requests.Session()
    base = {"sortBy": "popularity", "includeUnrunnableActors": "true"}
    for extra in ({"limit": 1}, {"limit": 100}, {"limit": 1000}, {"limit": 1000, "offset": 1000},
                  {"limit": 1000, "category": "SOCIAL_MEDIA"}):
        d = get_json(s, args.api_url, {**base, **extra}) or {}
        meta = {k: d.get(k) for k in ("total", "count", "offset", "limit")}
        print(f"{extra} -> получено {len(d.get('items') or [])}, мета: {meta}")


def probe_ids_cmd(args) -> None:
    """Проверка источников ID на живой сети: Store API, сайт /store/categories, /v2/acts/{user}~{name}."""
    import requests
    from .apify_api import get_json
    s = requests.Session()
    data = get_json(s, args.api_url, {"limit": 3, "offset": 0, "sortBy": "popularity"})
    items = (data or {}).get("items") or []
    print("Store API: поля первого элемента:", sorted(items[0]) if items else "нет данных")
    print("Store API: id в элементах:", [it.get("id") for it in items])
    page = ids._get_html(s, ids.CATEGORIES_URL)
    found = ids.ids_from_html(page)
    print(f"Сайт {ids.CATEGORIES_URL}: HTML {len(page)} симв., ссылок категорий {len(ids.category_links(page))}, "
          f"ID в JSON страницы: {len(found)}", list(found.items())[:3])
    if items:
        u, n = items[0].get("username"), items[0].get("name")
        print(f"/v2/acts/{u}~{n} -> id:", ids.lookup_id(s, u, n))


def rescore(args) -> None:
    """Пересчёт ниш/услуг/рейтингов по последнему снимку из MySQL (после смены весов, словарей, --niches-file)."""
    cfg = make_cfg(args)
    conn = db.connect(db.dsn_from_env(args.dsn))
    try:
        db.init_schema(conn)
        last = db.latest_run(conn)
        if not last:
            raise RuntimeError("В MySQL ещё нет снимков: сначала запустите сбор (или --demo 3000)")
        run_id, date = last
        log.info("Пересчёт снимка %s (run_id=%s)", date, run_id)
        score_and_publish(conn, run_id, db.load_snapshot(conn, run_id), cfg)
        if args.xlsx:
            export_xlsx(conn, args.xlsx, cfg)
    finally:
        conn.close()


def rewrite_cmd(args) -> None:
    """Отдельный запуск рерайта по последнему опубликованному снимку (одна итерация = окно --rewrite-window-hours)."""
    conn = db.connect(db.dsn_from_env(args.dsn))
    try:
        db.init_schema(conn)
        content.run_rewrite(conn, make_rewrite_cfg(args))
    finally:
        conn.close()


def flush_cmd(args) -> None:
    """Опубликовать всё, что осталось в очереди service_content_staging (например, после аварийной остановки)."""
    conn = db.connect(db.dsn_from_env(args.dsn))
    try:
        db.init_schema(conn)
        log.info("Опубликовано карточек: %s", content.ContentStore(conn).flush())
    finally:
        conn.close()


def probe_cmd(args) -> None:
    """Проверка одного Actor'а: что реально отдаёт API по описанию (запустите один раз на живой сети)."""
    import requests
    user, _, name = args.rewrite_probe.partition("/")
    rc = make_rewrite_cfg(args)
    d = content.fetch_details(requests.Session(), user, name, rc)
    print(json.dumps(d, ensure_ascii=False, indent=2) if d else "Actor не найден или API не вернул данных")


def export_xlsx(conn, path: Path, cfg: ScoreConfig) -> None:
    """Необязательная выгрузка витрины в Excel (данные берутся из тех же представлений, что и у фронта)."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    top = db.fetch_df(conn, "SELECT niche_name AS Niche, service_display_name AS Service, rank_no AS `Rank`, "
                            "actor_title AS Actor, developer AS Developer, ROUND(value_score,1) AS `Value score`, "
                            "ROUND(cost_1000,3) AS `Est. cost / 1000 (USD)`, pricing_summary AS Pricing, "
                            "total_users AS Users, users_30d AS `Users 30d`, rating AS `Avg rating`, "
                            "review_count AS `Ratings`, score_scope AS Scope, actor_url AS URL "
                            "FROM v_service_top_latest ORDER BY niche_name, service_display_name, rank_no")
    niches = db.fetch_df(conn, "SELECT niche_name AS Niche, niche_kind AS Kind, actors AS Actors, "
                               "services_count AS Services, eligible AS Eligible, total_users AS `Total users`, "
                               "ROUND(median_cost_1000,3) AS `Median cost / 1000`, "
                               "ROUND(min_cost_1000,3) AS `Min cost / 1000`, ROUND(median_rating,2) AS `Median rating` "
                               "FROM v_niche_overview_latest ORDER BY total_users DESC")
    services = db.fetch_df(conn, "SELECT niche_name AS Niche, service_display_name AS Service, service_kind AS Kind, "
                                 "actors AS Actors, eligible AS Eligible, total_users AS `Total users`, "
                                 "ROUND(median_cost_1000,3) AS `Median cost / 1000`, best_actor_title AS `Best actor`, "
                                 "ROUND(best_value_score,1) AS `Best score` "
                                 "FROM v_service_overview_latest ORDER BY niche_name, total_users DESC")
    formula = pd.DataFrame(formula_text(cfg), columns=["Элемент", "Как считается"])
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in [("Top by service", top), ("Niches", niches), ("Services", services), ("Value formula", formula)]:
            df.to_excel(xw, sheet_name=name, index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font, cell.fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
            for i, col in enumerate(ws.columns, 1):
                w = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:300])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 60)
    log.info("Excel: %s", path.resolve())


def seconds_until(hhmm: str) -> float:
    h, m = map(int, hhmm.split(":"))
    now = dt.datetime.now()
    nxt = now.replace(hour=h, minute=m, second=0, microsecond=0)
    if nxt <= now:
        nxt += dt.timedelta(days=1)
    return (nxt - now).total_seconds()


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="python -m tracker", description=__doc__)
    p.add_argument("--dsn", default=None, help="mysql://user:pass@host:3306/db (или переменная MYSQL_DSN)")
    p.add_argument("--mode", choices=["full", "top"], default="full",
                   help="full: максимальное покрытие (много запросов), top: быстрый срез по популярности")
    p.add_argument("--max-actors", type=int, default=DEFAULT_MAX_ACTORS, help="остановиться, набрав столько (0 = без лимита)")
    p.add_argument("--workers", type=int, default=6, help="число параллельных запросов")
    p.add_argument("--delay", type=float, default=0.2, help="пауза после каждого запроса в потоке, сек")
    p.add_argument("--max-terms", type=int, default=3000, help="максимум поисковых слов на фазе search-обхода")
    p.add_argument("--pricing-details", choices=["auto", "always", "never"], default="auto",
                   help="докачка деталей цены из /acts/{user}~{name}")
    p.add_argument("--exclude-unrunnable", action="store_true", help="не передавать includeUnrunnableActors=true")
    p.add_argument("--keep-days", type=int, default=365, help="хранить историю N дней (0 = вечно)")
    p.add_argument("--xlsx", type=Path, default=None, help="дополнительно выгрузить витрину в Excel")
    g = p.add_argument_group("ниши, услуги и рейтинг выгодности")
    g.add_argument("--per-service", type=int, default=3, help="сколько лучших Actor'ов в каждой услуге")
    g.add_argument("--min-niche-actors", type=int, default=5, help="минимум Actor'ов в нише")
    g.add_argument("--min-service-actors", type=int, default=3, help="минимум Actor'ов, чтобы услуга была самостоятельной")
    g.add_argument("--min-service-pool", type=int, default=5,
                   help="если подходящих Actor'ов в услуге меньше, шкалы считаются по всей нише")
    g.add_argument("--min-discovered-service", type=int, default=8,
                   help="порог автоподбора услуг: сколько Actor'ов ниши должны содержать слово")
    g.add_argument("--no-auto-services", action="store_true", help="только встроенный справочник услуг")
    g.add_argument("--min-users", type=int, default=10, help="минимум пользователей, чтобы попасть в рейтинг")
    g.add_argument("--weights", default="45,20,20,15", help="веса: цена,пользователи,оценка,число оценок")
    g.add_argument("--compute-cost", type=float, default=0.5, help="допущение: $ compute за 1000 результатов")
    g.add_argument("--monthly-results", type=int, default=10_000, help="допущение: результатов в месяц для аренды")
    g.add_argument("--bayes-prior", type=float, default=5.0, help="вес среднего по рынку в байесовской оценке")
    g.add_argument("--niches-file", type=Path, default=None, help="JSON с правками ниш и услуг (см. README)")
    r = p.add_argument_group("рерайт услуг для продажи (tracker/content.py)")
    r.add_argument("--rewrite", action="store_true", help="одна итерация рерайта по последнему снимку (без сбора)")
    r.add_argument("--with-rewrite", action="store_true", help="после ежедневного сбора запускать рерайт в том же прогоне")
    r.add_argument("--flush-content", action="store_true", help="опубликовать накопленную очередь рерайтов и выйти")
    r.add_argument("--catalog", action="store_true", help="одна итерация каталога: рерайт только НОВЫХ Apify ID")
    r.add_argument("--with-catalog", action="store_true", help="после ежедневного сбора запускать каталог в том же прогоне")
    r.add_argument("--catalog-scope", choices=["top", "eligible", "all"], default="top",
                   help="top: акторы из топов услуг, eligible: все с ценой и min-users, all: вообще все (много LLM-вызовов)")
    r.add_argument("--rewrite-probe", metavar="USER/NAME", default=None, help="показать, что API отдаёт по описанию Actor'а")
    r.add_argument("--rewrite-provider", choices=["llm", "template"], default="llm",
                   help="llm: ANTHROPIC_API_KEY, template: без ключа (для проверки конвейера)")
    r.add_argument("--rewrite-model", default=None, help="модель (или переменная REWRITE_MODEL)")
    r.add_argument("--rewrite-lang", default="en", help="язык карточек")
    r.add_argument("--rewrite-window-hours", type=float, default=3.0, help="длина одной итерации рерайта, часов")
    r.add_argument("--flush-minutes", type=float, default=15.0, help="как часто публиковать накопленное, минут")
    r.add_argument("--rewrite-workers", type=int, default=3, help="параллельных рерайтов")
    r.add_argument("--rewrite-sources", type=int, default=5, help="Actor'ов услуги в исходниках")
    r.add_argument("--rewrite-min-age-days", type=float, default=7, help="не переписывать услугу чаще, чем раз в N дней")
    r.add_argument("--rewrite-max-sim", type=float, default=0.20, help="макс. совпадение 4-грамм с исходниками")
    r.add_argument("--rewrite-niche", default=None, help="только ниша (slug), например instagram")
    r.add_argument("--rewrite-limit", type=int, default=0, help="максимум услуг за итерацию (0 = без лимита)")
    r.add_argument("--rewrite-force", action="store_true", help="переписать всё, игнорируя хэши и возраст")
    r.add_argument("--details-ttl-days", type=int, default=14, help="срок годности кэша описаний Actor'ов")
    r.add_argument("--no-fetch-details", action="store_true", help="не ходить в Apify за описаниями, брать из кэша")
    m = p.add_argument_group("мониторинг цен и ID акторов (tracker/monitor.py, tracker/ids.py)")
    m.add_argument("--monitor-prices", action="store_true", help="один проход мониторинга цен всех акторов (без рейтингов)")
    m.add_argument("--monitor-loop-hours", type=float, default=0, metavar="H",
                   help="с --monitor-prices: повторять проход каждые H часов (0 = один раз)")
    m.add_argument("--probe-api", action="store_true", help="показать реальные лимиты Store API (limit/offset/total)")
    m.add_argument("--probe-ids", action="store_true", help="проверить источники ID акторов (API, сайт, /v2/acts)")
    m.add_argument("--no-site-ids", action="store_true", help="не обходить apify.com/store/categories, только API")
    p.add_argument("--api-url", default=API_URL)
    p.add_argument("--daemon", action="store_true", help="работать постоянно, запуск раз в день")
    p.add_argument("--at", default="03:00", help="время ежедневного запуска HH:MM (локальное)")
    p.add_argument("--init-db", action="store_true", help="только создать/обновить схему и представления")
    p.add_argument("--rescore", action="store_true", help="пересчитать ниши/услуги/рейтинги по последнему снимку в MySQL")
    p.add_argument("--demo", type=int, default=0, metavar="N", help="вместо обхода Apify сгенерировать N тестовых Actor'ов")
    return p


def main() -> int:
    args = build_parser().parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    if args.init_db:
        conn = db.connect(db.dsn_from_env(args.dsn))
        db.init_schema(conn)
        conn.close()
        return 0
    if args.rescore:
        rescore(args)
        return 0
    if args.probe_api:
        probe_api_cmd(args)
        return 0
    if args.probe_ids:
        probe_ids_cmd(args)
        return 0
    if args.rewrite_probe:
        probe_cmd(args)
        return 0
    if args.monitor_prices:
        monitor_loop(args) if args.monitor_loop_hours else monitor_once(args)
        return 0
    if args.catalog:
        catalog_cmd(args)
        return 0
    if args.flush_content:
        flush_cmd(args)
        return 0
    if args.rewrite:
        rewrite_cmd(args)
        return 0
    if not args.daemon:
        run_once(args)
        return 0

    log.info("Режим демона: ежедневный запуск в %s", args.at)
    while True:
        wait = seconds_until(args.at)
        log.info("Следующий запуск через %.1f ч", wait / 3600)
        time.sleep(wait)
        try:
            run_once(args)
        except Exception:
            log.exception("Прогон завершился ошибкой, повторим завтра")


if __name__ == "__main__":
    sys.exit(main())
